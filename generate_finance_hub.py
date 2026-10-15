import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_finance_hub():
    wb = openpyxl.Workbook()
    
    # ----------------------------------------------------
    # ESTILOS COMUNES
    # ----------------------------------------------------
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E1B4B")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="6B7280")
    kpi_title_font = Font(name="Calibri", size=10, bold=True, color="4B5563")
    kpi_num_font = Font(name="Calibri", size=18, bold=True, color="1E1B4B")
    bold_font = Font(name="Calibri", size=11, bold=True)
    
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid") # Indigo
    header_fill_green = PatternFill(start_color="059669", end_color="059669", fill_type="solid") # Emerald
    header_fill_red = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid") # Red
    kpi_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    total_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    total_border = Border(
        top=Side(style='thin', color='111827'),
        bottom=Side(style='double', color='111827')
    )

    # ----------------------------------------------------
    # HOJA 1: DASHBOARD ANUAL
    # ----------------------------------------------------
    ws_dash = wb.active
    ws_dash.title = "📊 Dashboard Anual"
    ws_dash.views.sheetView[0].showGridLines = True

    ws_dash["B2"] = "FREELANCE FINANCE HUB 2026"
    ws_dash["B2"].font = title_font
    ws_dash["B3"] = "Resumen financiero automatizado para profesionales independientes y autónomos"
    ws_dash["B3"].font = subtitle_font

    # Tarjetas KPI
    kpis = [
        ("B5", "C6", "INGRESOS TOTALES (BASE)", "=SUM('💰 Ingresos'!E5:E100)"),
        ("E5", "F6", "GASTOS TOTALES (BASE)", "=SUM('📉 Gastos'!E5:E100)"),
        ("H5", "I6", "BENEFICIO NETO ANTES IMPUESTOS", "=B6-E6"),
        ("K5", "L6", "ESTIMACIÓN IVA A PAGAR (HACIENDA)", "=SUM('💰 Ingresos'!G5:G100)-SUM('📉 Gastos'!F5:F100)")
    ]

    for start_cell, end_cell, label, formula in kpis:
        c_label = ws_dash[start_cell]
        c_val = ws_dash[start_cell[0] + str(int(start_cell[1]) + 1)]
        c_label.value = label
        c_label.font = kpi_title_font
        c_val.value = formula
        c_val.font = kpi_num_font
        c_val.number_format = "#,##0.00 €"
        
        # Merge de 2 columnas para el KPI
        ws_dash.merge_cells(f"{start_cell}:{chr(ord(start_cell[0])+1)}{start_cell[1]}")
        ws_dash.merge_cells(f"{start_cell[0]}{int(start_cell[1])+1}:{chr(ord(start_cell[0])+1)}{int(start_cell[1])+1}")

    # Tabla resumen mensual
    ws_dash["B9"] = "Mes"
    ws_dash["C9"] = "Ingresos Brutos"
    ws_dash["D9"] = "Gastos Totales"
    ws_dash["E9"] = "Beneficio Neto"
    ws_dash["F9"] = "IVA Repercutido"
    ws_dash["G9"] = "IVA Soportado"

    for col in ["B", "C", "D", "E", "F", "G"]:
        c = ws_dash[f"{col}9"]
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if col=="B" else "right")

    meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
    for i, m in enumerate(meses, start=10):
        ws_dash[f"B{i}"] = m
        ws_dash[f"C{i}"] = f"=SUMIFS('💰 Ingresos'!E:E, '💰 Ingresos'!A:A, \">=\"&DATE(2026,{i-9},1), '💰 Ingresos'!A:A, \"<=\"&EOMONTH(DATE(2026,{i-9},1),0))"
        ws_dash[f"D{i}"] = f"=SUMIFS('📉 Gastos'!E:E, '📉 Gastos'!A:A, \">=\"&DATE(2026,{i-9},1), '📉 Gastos'!A:A, \"<=\"&EOMONTH(DATE(2026,{i-9},1),0))"
        ws_dash[f"E{i}"] = f"=C{i}-D{i}"
        ws_dash[f"F{i}"] = f"=SUMIFS('💰 Ingresos'!G:G, '💰 Ingresos'!A:A, \">=\"&DATE(2026,{i-9},1), '💰 Ingresos'!A:A, \"<=\"&EOMONTH(DATE(2026,{i-9},1),0))"
        ws_dash[f"G{i}"] = f"=SUMIFS('📉 Gastos'!F:F, '📉 Gastos'!A:A, \">=\"&DATE(2026,{i-9},1), '📉 Gastos'!A:A, \"<=\"&EOMONTH(DATE(2026,{i-9},1),0))"
        
        for col in ["C", "D", "E", "F", "G"]:
            ws_dash[f"{col}{i}"].number_format = "#,##0.00 €"
            ws_dash[f"{col}{i}"].border = thin_border

    # Fila de Totales
    ws_dash["B22"] = "TOTAL ANUAL"
    ws_dash["B22"].font = bold_font
    for col in ["C", "D", "E", "F", "G"]:
        c = ws_dash[f"{col}22"]
        c.value = f"=SUM({col}10:{col}21)"
        c.font = bold_font
        c.number_format = "#,##0.00 €"
        c.border = total_border
        c.fill = total_fill

    # ----------------------------------------------------
    # HOJA 2: REGISTRO DE INGRESOS
    # ----------------------------------------------------
    ws_inc = wb.create_sheet(title="💰 Ingresos")
    ws_inc.views.sheetView[0].showGridLines = True

    inc_headers = ["Fecha", "Nº Factura", "Cliente", "Concepto", "Base Imponible (€)", "IVA %", "Cuota IVA (€)", "IRPF %", "Retención IRPF (€)", "Total Cobro (€)", "Estado"]
    for c_idx, h in enumerate(inc_headers, start=1):
        cell = ws_inc.cell(row=4, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill_green
        cell.alignment = Alignment(horizontal="center" if "Fecha" in h or "%" in h or "Estado" in h else "left")

    # Filas de ejemplo realistas
    sample_inc = [
        ("2026-01-15", "FAC-001", "Acme Corp S.L.", "Diseño interfaz web SaaS", 1800, 21, "=E5*(F5/100)", 15, "=E5*(H5/100)", "=E5+G5-I5", "Cobrado"),
        ("2026-01-28", "FAC-002", "Innovatech Studio", "Desarrollo frontend React", 2400, 21, "=E6*(F6/100)", 15, "=E6*(H6/100)", "=E6+G6-I6", "Cobrado"),
        ("2026-02-10", "FAC-003", "Global Marketing S.A.", "Auditoría SEO técnica", 950, 21, "=E7*(F7/100)", 15, "=E7*(H7/100)", "=E7+G7-I7", "Pendiente"),
        ("2026-02-25", "FAC-004", "Restaurante El Rincón", "Mantenimiento web trimestral", 450, 21, "=E8*(F8/100)", 15, "=E8*(H8/100)", "=E8+G8-I8", "Cobrado")
    ]

    for r_idx, row_data in enumerate(sample_inc, start=5):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws_inc.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            if c_idx in [5, 7, 9, 10]:
                cell.number_format = "#,##0.00 €"
            elif c_idx in [6, 8]:
                cell.number_format = "0\"%\""

    # ----------------------------------------------------
    # HOJA 3: REGISTRO DE GASTOS
    # ----------------------------------------------------
    ws_exp = wb.create_sheet(title="📉 Gastos")
    ws_exp.views.sheetView[0].showGridLines = True

    exp_headers = ["Fecha", "Proveedor / Empresa", "Categoría", "Descripción / Factura", "Base Imponible (€)", "IVA Soportado (€)", "Total Pagado (€)"]
    for c_idx, h in enumerate(exp_headers, start=1):
        cell = ws_exp.cell(row=4, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill_red
        cell.alignment = Alignment(horizontal="center" if "Fecha" in h else "left")

    sample_exp = [
        ("2026-01-05", "Seguridad Social", "Cuota Autónomos", "Cuota mensual autónomos RETA", 295.00, 0.00, "=E5+F5"),
        ("2026-01-10", "Gestoría Online S.L.", "Gestoría", "Asesoría fiscal y laboral mensual", 60.00, 12.60, "=E6+F6"),
        ("2026-01-12", "Figma Inc", "Software", "Suscripción anual Professional", 144.00, 30.24, "=E7+F7"),
        ("2026-01-15", "Google Cloud / Workspace", "Software", "Email corporativo y almacenamiento", 12.00, 2.52, "=E8+F8"),
        ("2026-02-05", "Seguridad Social", "Cuota Autónomos", "Cuota mensual autónomos RETA", 295.00, 0.00, "=E9+F9"),
        ("2026-02-10", "Gestoría Online S.L.", "Gestoría", "Asesoría fiscal mensual", 60.00, 12.60, "=E10+F10"),
    ]

    for r_idx, row_data in enumerate(sample_exp, start=5):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws_exp.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            if c_idx in [5, 6, 7]:
                cell.number_format = "#,##0.00 €"

    # Autoajuste de columnas
    for sheet in [ws_dash, ws_inc, ws_exp]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output_path = r"C:\Users\Miguel\.gemini\antigravity\scratch\freelance-toolkit-saas\entregables\Freelance_Finance_Hub_2026.xlsx"
    wb.save(output_path)
    print(f"Creado Finance Hub en {output_path}")

create_finance_hub()
